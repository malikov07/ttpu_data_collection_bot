"""Certificates and awards: students send them in the bot, staff accept or reject."""

from __future__ import annotations

import pytest
from aiogram import methods
from sqlalchemy import select

from app.callbacks import CertCb, CertReviewCb, StaffCb
from app.db import repo
from app.db.models import AuditLog, Certificate, CertStatus, CertType, StaffRole
from app.services.certificates import normalize_result
from app.services.validators import ValidationError
from tests.test_flows import login_bot, register, setup_groups, student



@pytest.mark.parametrize(
    "cert_type, raw, stored",
    [
        (CertType.IELTS, "7", "7.0"),
        (CertType.IELTS, "6,5", "6.5"),
        (CertType.TOEFL, "95", "95"),
        (CertType.SAT, "1450", "1450"),
        (CertType.DUOLINGO, "125", "125"),
        (CertType.CEFR, "b2", "B2"),
        (CertType.NATIONAL, "  Matematika   A+ ", "Matematika A+"),
    ],
)
def test_results_are_normalized(cert_type, raw, stored):
    assert normalize_result(cert_type, raw) == stored


@pytest.mark.parametrize(
    "cert_type, raw",
    [
        (CertType.IELTS, "7.3"),
        (CertType.IELTS, "9.5"),
        (CertType.IELTS, "seven"),
        (CertType.TOEFL, "121"),
        (CertType.TOEFL, "95.5"),
        (CertType.SAT, "1455"),
        (CertType.SAT, "300"),
        (CertType.DUOLINGO, "121"),
        (CertType.CEFR, "B3"),
        (CertType.OTHER, "A"),
        (CertType.OLYMPIAD, "1234"),
    ],
)
def test_invalid_results_are_rejected(cert_type, raw):
    with pytest.raises(ValidationError):
        normalize_result(cert_type, raw)


async def certificates(h, uid) -> list[Certificate]:
    async with h.sessions() as s:
        return (await repo.get_student_by_tg(s, uid)).certificates


async def add_ielts(h, u, score="7.5", file="ielts.pdf"):
    await h.text(u, "🏆 My certificates")
    await h.click(u, CertCb(action="add").pack())
    await h.click(u, CertCb(action="type", value="ielts").pack())
    await h.text(u, score)
    await h.document(u, file, "application/pdf")
    await h.click(u, CertCb(action="submit").pack())


async def test_student_adds_certificates(h):
    await setup_groups(h)
    u = await register(h, uid=100)
    # After registering, the menu has the certificates button.
    menu = h.tg.last(methods.SendMessage).reply_markup
    assert "🏆 My certificates" in [b.text for row in menu.keyboard for b in row]

    await h.text(u, "🏆 My certificates")
    assert "haven't added any" in h.tg.texts(100)[-1]
    await h.click(u, CertCb(action="add").pack())
    await h.click(u, CertCb(action="type", value="ielts").pack())
    assert "overall band score" in h.tg.texts(100)[-1]
    await h.text(u, "7.3")
    assert "steps of 0.5" in h.tg.texts(100)[-1]
    await h.text(u, "7,5")
    assert "IELTS 7.5" in h.tg.texts(100)[-1] and "PDF" in h.tg.texts(100)[-1]
    await h.document(u, "cv.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "PDF or a photo" in h.tg.texts(100)[-1]
    await h.document(u, "ielts.pdf", "application/pdf")
    assert "Send it for review?" in h.tg.texts(100)[-1]
    await h.click(u, CertCb(action="submit").pack())
    assert any("sent for review" in t for t in h.tg.texts(100))
    assert "IELTS 7.5</b> · ⏳ under review" in h.tg.texts(100)[-1]

    # Photos of pages, and a level chosen with a button.
    await h.click(u, CertCb(action="add").pack())
    await h.click(u, CertCb(action="type", value="cefr").pack())
    await h.click(u, CertCb(action="result", value="B2").pack())
    await h.photo(u, "cefr-1")
    await h.photo(u, "cefr-2")
    await h.document(u, "late.pdf", "application/pdf")
    assert "either one PDF or photos" in h.tg.texts(100)[-1]
    await h.click(u, CertCb(action="done").pack())
    assert "Files: 2" in h.tg.texts(100)[-1]
    await h.click(u, CertCb(action="submit").pack())

    items = await certificates(h, 100)
    assert [(c.type, c.result, c.status) for c in items] == [
        (CertType.IELTS, "7.5", CertStatus.PENDING),
        (CertType.CEFR, "B2", CertStatus.PENDING),
    ]
    assert items[0].files[0]["name"] == "ielts.pdf"
    assert [f["file_id"] for f in items[1].files] == ["cefr-1", "cefr-2"]
    async with h.sessions() as s:
        actions = [a.action for a in (await s.scalars(select(AuditLog).where(AuditLog.action.like("certificate.%")))).all()]
    assert actions == ["certificate.submit", "certificate.submit"]


async def test_only_registered_students_add_certificates(h):
    await setup_groups(h)
    u = h.user(120, "New")
    await h.text(u, "/start")
    await h.click(u, "lang:en")
    await h.text(u, "🏆 My certificates")
    assert "Agree and start" in [text for text, _ in h.buttons()][0]
    await h.click(u, CertCb(action="add").pack())
    await h.click(u, CertCb(action="type", value="ielts").pack())
    async with h.sessions() as s:
        assert (await s.scalars(select(Certificate))).all() == []


async def test_student_removes_a_certificate_until_it_is_accepted(h):
    await setup_groups(h)
    u = await register(h, uid=100)
    await add_ielts(h, u)
    await add_ielts(h, u, score="8", file="ielts2.pdf")
    first, second = await certificates(h, 100)

    await h.click(u, CertCb(action="del", id=first.id).pack())
    assert "Remove <b>IELTS 7.5</b>?" in h.tg.texts(100)[-1]
    await h.click(u, CertCb(action="del_ok", id=first.id).pack())
    assert [c.id for c in await certificates(h, 100)] == [second.id]

    async with h.sessions() as s:
        cert = await s.get(Certificate, second.id)
        cert.status = CertStatus.APPROVED
        await s.commit()
    await h.click(u, CertCb(action="del_ok", id=second.id).pack())
    assert "can't be removed" in h.tg.last(methods.AnswerCallbackQuery).text
    assert len(await certificates(h, 100)) == 1

    # Someone else's certificate can't be removed either.
    other = await register(h, uid=101, number="AB7654321", personal="32103059999999")
    await h.click(other, CertCb(action="del_ok", id=second.id).pack())
    assert len(await certificates(h, 100)) == 1


async def test_leader_reviews_certificates(h):
    await setup_groups(h)
    await h.create_account("leader1", StaffRole.LEADER, "SE-24-01")
    await h.create_account("leader2", StaffRole.LEADER, "SE-24-02")
    leader = await login_bot(h, 800, "leader1")
    other_leader = await login_bot(h, 801, "leader2")
    u = await register(h, uid=100, group="SE-24-01")

    await h.text(leader, "/certificates")
    assert "No certificates are waiting" in h.tg.texts(800)[-1]
    await add_ielts(h, u)
    assert any("Aliyev Vali Karimovich sent a certificate: IELTS 7.5" in t for t in h.tg.texts(800))  # leader notified
    cert = (await certificates(h, 100))[0]

    await h.text(leader, "/certificates")
    assert "Awaiting review:</b> 1" in h.tg.texts(800)[-1]
    assert "⏳ IELTS 7.5 · Aliyev Vali Karimovich · SE-24-01" in [text for text, _ in h.buttons()]

    # The student card links to the certificates.
    s = await student(h, 100)
    await h.click(leader, StaffCb(action="card", student_id=s.id).pack())
    assert "Certificates: 1 · ⏳ 1 under review" in h.tg.last(methods.SendPhoto).caption
    await h.click(leader, h.button("Certificates (1)"))
    await h.click(leader, h.button("IELTS 7.5"))
    assert h.tg.last(methods.SendDocument).document == "doc-ielts.pdf"
    assert "under review" in h.tg.texts(800)[-1]

    await h.click(other_leader, CertReviewCb(action="open", id=cert.id).pack())
    assert h.tg.last(methods.AnswerCallbackQuery).show_alert  # another group's student
    await h.click(other_leader, CertReviewCb(action="ok", id=cert.id).pack())
    assert (await certificates(h, 100))[0].status == CertStatus.PENDING

    await h.click(leader, CertReviewCb(action="no", id=cert.id).pack())
    assert "Type the reason" in h.tg.texts(800)[-1]
    await h.text(leader, "The photo is blurry,   please send a scan")
    cert = (await certificates(h, 100))[0]
    assert cert.status == CertStatus.REJECTED and cert.note == "The photo is blurry, please send a scan"
    assert cert.reviewed_by == "leader1"
    assert "not accepted: <b>IELTS 7.5</b>\nReason: <i>The photo is blurry, please send a scan</i>" in h.tg.texts(100)[-1]

    # A decision can be changed; accepting clears the reason.
    await h.click(leader, CertReviewCb(action="ok", id=cert.id).pack())
    cert = (await certificates(h, 100))[0]
    assert cert.status == CertStatus.APPROVED and cert.note is None
    assert "was accepted: <b>IELTS 7.5</b>" in h.tg.texts(100)[-1]
    await h.text(u, "🏆 My certificates")
    assert "✅ accepted" in h.tg.texts(100)[-1]
    await h.text(leader, "/certificates")
    assert "No certificates are waiting" in h.tg.texts(800)[-1]

    async with h.sessions() as db:
        actions = [a.action for a in (await db.scalars(select(AuditLog).where(AuditLog.action.like("certificate.%")))).all()]
    assert actions == ["certificate.submit", "certificate.view", "certificate.reject", "certificate.approve"]


async def test_reject_without_reason(h):
    await setup_groups(h)
    await h.create_account("tutor1", StaffRole.TUTOR)
    tutor = await login_bot(h, 700, "tutor1")
    u = await register(h, uid=100)
    await add_ielts(h, u)
    cert = (await certificates(h, 100))[0]
    await h.click(tutor, CertReviewCb(action="no", id=cert.id).pack())
    await h.click(tutor, CertReviewCb(action="skip", id=cert.id).pack())
    cert = (await certificates(h, 100))[0]
    assert cert.status == CertStatus.REJECTED and cert.note is None
    assert h.tg.texts(100)[-1].endswith("not accepted: <b>IELTS 7.5</b>")



async def test_certificates_are_sent_with_registration(h):
    from tests.conftest import passport_result
    from tests.test_flows import start_registration

    await setup_groups(h)
    h.vision.documents["p"] = passport_result()
    u = await start_registration(h, 100)
    await h.photo(u, "p")
    await h.click(u, "reg:correct")
    await h.contact(u, "998901234567")
    await h.text(u, "SE-24-01")
    await h.photo(u, "face", size=(600, 800))
    await h.document(u, "cv.pdf", "application/pdf")
    step = h.tg.texts(100)[-1]
    assert "Certificates" in step and "7/7" in step
    assert [text for text, _ in h.buttons()] == ["➕ Add a certificate", "I have none ›"]

    # Changing their mind halfway keeps the form; "Back" (not "Cancel") returns to the step.
    await h.click(u, CertCb(action="add").pack())
    assert ("‹ Back", CertCb(action="back").pack()) in h.buttons()
    await h.click(u, CertCb(action="type", value="sat").pack())
    await h.click(u, CertCb(action="back").pack())
    assert "Certificates" in h.tg.texts(100)[-1] and "Added" not in h.tg.texts(100)[-1]

    await h.click(u, CertCb(action="add").pack())
    await h.click(u, CertCb(action="type", value="ielts").pack())
    await h.text(u, "7")
    await h.document(u, "ielts.pdf", "application/pdf")
    assert "sent with your data" in h.tg.texts(100)[-1]
    await h.click(u, CertCb(action="submit").pack())
    await h.click(u, CertCb(action="add").pack())
    await h.click(u, CertCb(action="type", value="national").pack())
    await h.text(u, "Matematika A+")
    await h.photo(u, "milliy")
    await h.click(u, CertCb(action="done").pack())
    await h.click(u, CertCb(action="submit").pack())
    await h.click(u, CertCb(action="add").pack())
    await h.click(u, CertCb(action="type", value="sat").pack())
    await h.text(u, "1400")
    await h.photo(u, "sat")
    await h.click(u, CertCb(action="done").pack())
    await h.click(u, CertCb(action="submit").pack())
    step = h.tg.texts(100)[-1]
    assert "✅ IELTS 7.0" in step and "✅ National certificate: Matematika A+" in step and "✅ SAT 1400" in step
    await h.click(u, h.button("🗑 SAT 1400"))  # added by mistake
    assert "SAT 1400" not in h.tg.texts(100)[-1]
    async with h.sessions() as s:
        assert (await s.scalars(select(Certificate))).all() == []  # nothing is saved before «Submit»

    await h.click(u, h.button("Continue"))
    review = h.tg.last(methods.SendPhoto).caption
    assert "🏆 IELTS 7.0, National certificate: Matematika A+" in review
    # The review can go back to the certificates.
    await h.click(u, "reg:edit")
    await h.click(u, h.button("🏆 Certificates"))
    await h.click(u, h.button("Continue"))
    await h.click(u, "reg:submit")

    items = await certificates(h, 100)
    assert [(c.type, c.result, c.status) for c in items] == [
        (CertType.IELTS, "7.0", CertStatus.PENDING),
        (CertType.NATIONAL, "Matematika A+", CertStatus.PENDING),
    ]
    assert items[1].files[0]["file_id"] == "milliy"
    await h.text(u, "🏆 My certificates")
    assert "IELTS 7.0</b> · ⏳ under review" in h.tg.texts(100)[-1]


async def test_registration_without_certificates_says_so(h):
    from tests.conftest import passport_result
    from tests.test_flows import start_registration

    await setup_groups(h)
    h.vision.documents["p"] = passport_result()
    u = await start_registration(h, 100)
    await h.photo(u, "p")
    await h.click(u, "reg:correct")
    await h.contact(u, "998901234567")
    await h.text(u, "SE-24-01")
    await h.photo(u, "face", size=(600, 800))
    await h.document(u, "cv.pdf", "application/pdf")
    await h.click(u, "reg:submit")  # the certificates step can't be skipped silently
    assert "Certificates" in h.tg.texts(100)[-1] and await student(h, 100) is None
    await h.click(u, "reg:certs_done")
    assert "🏆 no certificates" in h.tg.last(methods.SendPhoto).caption
    await h.click(u, "reg:submit")
    assert (await student(h, 100)) is not None and await certificates(h, 100) == []


# ----------------------------------------------------------------------------- website


async def test_website_lists_and_reviews_certificates(web, h, monkeypatch):
    import io

    from openpyxl import load_workbook

    from app.services import documents
    from tests.test_web import login_as, seed

    async def fake_download(bot, file_id: str) -> bytes:
        return b"%PDF-1.4 " + file_id.encode()

    monkeypatch.setattr(documents, "download", fake_download)
    g1, g2, s1, s2 = await seed(h)
    u1, u2 = h.user(100, "Vali", username="vali100"), h.user(101, "Dilnoza", username="vali101")
    await add_ielts(h, u1)
    await add_ielts(h, u2, score="6", file="dilnoza.pdf")
    (c1,), (c2,) = await certificates(h, 100), await certificates(h, 101)

    leader = await login_as(web, h, "leader1", StaffRole.LEADER, "SE-24-01")
    r = (await leader.get("/api/certificates")).json()
    assert r["total"] == 1 and r["counts"] == {"pending": 1, "approved": 0, "rejected": 0}
    item = r["items"][0]
    assert (item["type"], item["result"], item["status"], item["can_review"]) == ("ielts", "7.5", "pending", True)
    assert item["student"]["full_name"] == "Aliyev Vali Karimovich" and item["files"][0]["name"] == "ielts.pdf"
    # Another group's certificate: not found, not reviewable.
    assert (await leader.get(f"/api/certificates/{c2.id}/files/0")).status_code == 404
    assert (await leader.post(f"/api/certificates/{c2.id}/review", json={"status": "approved"})).status_code == 404

    f = await leader.get(f"/api/certificates/{c1.id}/files/0")
    assert f.content == b"%PDF-1.4 doc-ielts.pdf" and f.headers["content-type"] == "application/pdf"
    r = await leader.post(f"/api/certificates/{c1.id}/review", json={"status": "rejected", "note": "  Blurry   scan "})
    assert r.status_code == 200 and (r.json()["status"], r.json()["note"]) == ("rejected", "Blurry scan")
    assert "Reason: <i>Blurry scan</i>" in h.tg.texts(100)[-1]  # the student is told in the bot
    assert (await leader.post(f"/api/certificates/{c1.id}/review", json={"status": "maybe"})).status_code == 422

    admin = await login_as(web, h, "admin1", StaffRole.ADMIN)
    r = (await admin.get("/api/certificates", params={"status": "pending"})).json()
    assert [i["id"] for i in r["items"]] == [c2.id] and r["counts"]["rejected"] == 1
    assert (await admin.post(f"/api/certificates/{c2.id}/review", json={"status": "approved"})).json()["status"] == "approved"
    assert "was accepted: <b>IELTS 6.0</b>" in h.tg.texts(101)[-1]
    assert (await admin.get("/api/stats")).json()["certificates"] == {"pending": 0, "approved": 1}

    detail = (await admin.get(f"/api/students/{s2.id}")).json()
    assert [(c["result"], c["status"], c["reviewed_by"]) for c in detail["certificates"]] == [("6.0", "approved", "admin1")]
    assert {"certificate.submit", "certificate.approve"} <= {e["action"] for e in detail["history"]}

    ws = load_workbook(io.BytesIO((await admin.get("/api/export/students.xlsx?lang=en")).content)).active
    assert ws["R1"].value == "Certificates"
    assert {ws["R2"].value or "", ws["R3"].value or ""} == {"", "IELTS 6.0"}  # rejected ones are left out
