"""Website API: password sign-in, accounts, permissions per role, editing,
documents, export, settings, audit."""

from __future__ import annotations

import io

import pytest
from sqlalchemy import select

from app.db import repo
from app.db.models import AuditLog, StaffRole
from app.services import vision
from tests.test_flows import register, setup_groups

pytestmark = pytest.mark.asyncio
PDF = b"%PDF-1.4\n%fake\n"
JPEG = b"\xff\xd8\xff\xe0fake-jpeg"


async def seed(h):
    await setup_groups(h)
    await register(h, uid=100, group="SE-24-01")
    await register(h, uid=101, group="SE-24-02", number="AB7654321", personal="32103059999999", surname="KARIMOVA", given="DILNOZA", sex="F")
    async with h.sessions() as s:
        g1, g2 = await repo.get_group_by_name(s, "SE-24-01"), await repo.get_group_by_name(s, "SE-24-02")
        s1, s2 = await repo.get_student_by_tg(s, 100), await repo.get_student_by_tg(s, 101)
    return g1, g2, s1, s2


async def admin_client(web, h):
    await h.create_account("admin", StaffRole.ADMIN)
    c, r = await web.login("admin")
    assert r.status_code == 204, r.text
    return c


async def test_password_login(web, h):
    await h.create_account("admin", StaffRole.ADMIN)
    c = web.client()
    assert (await c.get("/api/me")).status_code == 401
    assert (await c.get("/api/config")).json()["bot_username"] == "ttpu_test_bot"

    _, r = await web.login("admin", "wrong")
    assert r.status_code == 401 and r.json()["detail"] == "invalid_credentials"
    _, r = await web.login("nobody", "Secret123")
    assert r.status_code == 401 and r.json()["detail"] == "invalid_credentials"  # same answer: no user enumeration

    c, r = await web.login("ADMIN")
    assert r.status_code == 204
    cookie = r.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=lax" in cookie
    me = (await c.get("/api/me")).json()
    assert me["role"] == "admin" and me["username"] == "admin" and not me["must_change_password"]

    assert (await c.post("/api/auth/logout", headers={"X-Requested-With": ""})).status_code == 403  # CSRF
    assert (await c.post("/api/auth/logout")).status_code == 204
    assert (await c.get("/api/me")).status_code == 401

    async with h.sessions() as s:
        actions = [a.action for a in (await s.scalars(select(AuditLog))).all()]
    assert "auth.login_failed" in actions and "auth.login" in actions


async def test_login_rate_limit(web, h):
    await h.create_account("admin", StaffRole.ADMIN)
    codes = [(await web.login("admin", "nope"))[1].status_code for _ in range(6)]
    assert codes[:5] == [401] * 5 and codes[5] == 429
    assert (await web.login("admin"))[1].status_code == 429  # locked even with the right password


async def test_temporary_password_must_be_changed(web, h):
    admin = await admin_client(web, h)
    await setup_groups(h)
    r = await admin.post("/api/accounts", json={"username": "tutor.one", "display_name": "Dilnoza", "roles": [{"role": "tutor"}]})
    assert r.status_code == 201
    temp = r.json()["temporary_password"]
    assert len(temp) == 12 and r.json()["must_change_password"]

    tutor, resp = await web.login("tutor.one", temp)
    assert resp.status_code == 204
    assert (await tutor.get("/api/me")).json()["must_change_password"]
    assert (await tutor.get("/api/students")).json()["detail"] == "password_change_required"

    r = await tutor.post("/api/auth/password", json={"current_password": temp, "new_password": "short"})
    assert r.status_code == 422 and r.json()["detail"]["error"] == "password_too_short"
    r = await tutor.post("/api/auth/password", json={"current_password": "bad", "new_password": "NewPass2026"})
    assert r.json()["detail"]["error"] == "wrong_password"
    assert (await tutor.post("/api/auth/password", json={"current_password": temp, "new_password": "NewPass2026"})).status_code == 204
    assert (await tutor.get("/api/students")).status_code == 200

    # admin resets it: the tutor is signed out everywhere
    account_id = (await tutor.get("/api/me")).json()["id"]
    new_temp = (await admin.post(f"/api/accounts/{account_id}/reset-password")).json()["temporary_password"]
    assert (await tutor.get("/api/me")).status_code == 401
    assert (await web.login("tutor.one", new_temp))[1].status_code == 204


async def test_account_management_guards(web, h):
    admin = await admin_client(web, h)
    await setup_groups(h)
    me_id = (await admin.get("/api/me")).json()["id"]
    assert (await admin.post("/api/accounts", json={"username": "Bad Name", "roles": [{"role": "tutor"}]})).json()["detail"]["error"] == "invalid_username"
    assert (await admin.post("/api/accounts", json={"username": "admin", "roles": [{"role": "tutor"}]})).status_code == 409
    r = await admin.post("/api/accounts", json={"username": "lead", "roles": [{"role": "leader"}]})
    assert r.json()["detail"]["error"] == "group_required"
    assert (await admin.patch(f"/api/accounts/{me_id}", json={"is_active": False})).json()["detail"] == "cannot_change_self"
    assert (await admin.delete(f"/api/accounts/{me_id}")).status_code == 409

    async with h.sessions() as s:
        g = await repo.get_group_by_name(s, "SE-24-01")
    r = await admin.post("/api/accounts", json={"username": "lead", "roles": [{"role": "leader", "group_id": g.id}]})
    lead = r.json()
    leader, _ = await web.login("lead", lead["temporary_password"])
    await leader.post("/api/auth/password", json={"current_password": lead["temporary_password"], "new_password": "Leader2026"})
    assert (await leader.get("/api/me")).status_code == 200
    assert (await admin.patch(f"/api/accounts/{lead['id']}", json={"is_active": False})).json()["is_active"] is False
    assert (await leader.get("/api/me")).status_code == 401  # disabled: signed out immediately
    assert (await admin.delete(f"/api/accounts/{lead['id']}")).status_code == 204


async def login_as(web, h, username, role, group=None):
    await h.create_account(username, role, group)
    c, r = await web.login(username)
    assert r.status_code == 204
    return c


async def test_leader_sees_and_edits_only_own_group(web, h):
    g1, g2, s1, s2 = await seed(h)
    leader = await login_as(web, h, "leader1", StaffRole.LEADER, "SE-24-01")

    listing = (await leader.get("/api/students")).json()
    assert [s["full_name"] for s in listing["items"]] == ["Aliyev Vali Karimovich"]
    item = listing["items"][0]
    assert item["document"] == {"type": "passport", "number": "AB1234567", "expiry": "2031-05-20", "pinfl": "32103051234567", "nationality": "UZB"}
    assert (await leader.get(f"/api/students/{s2.id}")).status_code == 404
    assert [g["name"] for g in (await leader.get("/api/groups")).json()] == ["SE-24-01"]

    r = await leader.patch(f"/api/students/{s1.id}", json={"phone": "+998 91 111 22 33", "middle_name": "karim o'g'li"})
    assert r.status_code == 200, r.text
    assert r.json()["phone"] == "+998911112233" and r.json()["full_name"] == "Aliyev Vali Karim oʻgʻli"
    r = await leader.patch(f"/api/students/{s1.id}", json={"pinfl": "123"})
    assert r.json()["detail"]["field"] == "pinfl"
    r = await leader.patch(f"/api/students/{s1.id}", json={"group_id": g2.id})
    assert r.json()["detail"]["field"] == "group_id"
    assert (await leader.patch(f"/api/students/{s2.id}", json={"phone": "901234567"})).status_code == 404
    assert (await leader.delete(f"/api/students/{s1.id}")).status_code == 403
    for path in ("/api/accounts", "/api/settings", "/api/audit"):
        assert (await leader.get(path)).status_code == 403, path


async def test_tutor_and_admin(web, h):
    g1, g2, s1, s2 = await seed(h)
    tutor = await login_as(web, h, "tutor1", StaffRole.TUTOR)
    assert (await tutor.get("/api/students")).json()["total"] == 2
    assert (await tutor.get("/api/students?q=32103059999999")).json()["total"] == 1  # PINFL search
    assert (await tutor.get("/api/students?q=ab76")).json()["total"] == 1  # document number
    r = await tutor.patch(f"/api/students/{s2.id}", json={"group_id": g1.id, "doc_number": "ab 7654322"})
    assert r.json()["group"]["name"] == "SE-24-01" and r.json()["document"]["number"] == "AB7654322"
    assert (await tutor.delete(f"/api/students/{s2.id}")).status_code == 403

    admin = await login_as(web, h, "admin1", StaffRole.ADMIN)
    assert (await admin.delete(f"/api/students/{s2.id}")).status_code == 204
    detail = (await admin.get(f"/api/students/{s1.id}")).json()
    assert detail["can_delete"] and [e["action"] for e in detail["history"]] == ["student.register"]


async def test_documents(web, h, monkeypatch):
    g1, g2, s1, s2 = await seed(h)
    admin = await login_as(web, h, "admin1", StaffRole.ADMIN)

    r = await admin.post(f"/api/students/{s1.id}/documents/cv", files=[("files", ("new-cv.pdf", PDF, "application/pdf"))])
    assert r.status_code == 200, r.text
    assert r.json()["documents"]["cv"]["pages"][0]["source"] == "upload"
    f = await admin.get(f"/api/students/{s1.id}/documents/cv/0")
    assert f.content == PDF and f.headers["cache-control"] == "private, no-store"

    r = await admin.post(f"/api/students/{s1.id}/documents/cv", files=[("files", ("cv.pdf", b"MZ\x90evil", "application/pdf"))])
    assert r.json()["detail"]["error"] == "unsupported_type"  # content, not the name, decides
    r = await admin.post(f"/api/students/{s1.id}/documents/passport", files=[("files", ("p.pdf", PDF, "application/pdf"))])
    assert r.json()["detail"]["error"] == "unsupported_type"  # passports must be images

    async def portrait_check(data: bytes):
        return "no_face" if data == JPEG else None

    monkeypatch.setattr(vision, "check_portrait", portrait_check)
    r = await admin.post(f"/api/students/{s1.id}/documents/photo", files=[("files", ("me.jpg", JPEG, "image/jpeg"))])
    assert r.json()["detail"]["error"] == "photo_no_face"

    detail = (await admin.get(f"/api/students/{s1.id}")).json()
    assert {"document.view", "document.replace"} <= {e["action"] for e in detail["history"]}


async def test_groups_settings_stats_export(web, h):
    g1, g2, s1, s2 = await seed(h)
    admin = await login_as(web, h, "admin1", StaffRole.ADMIN)
    r = await admin.post("/api/groups", json={"names": "ME-24-01, bad name!, SE-24-01"})
    assert r.json() == {"added": ["ME-24-01"], "existing": ["SE-24-01"], "invalid": ["BAD-NAME!"]}
    assert (await admin.patch(f"/api/groups/{g1.id}", json={"name": "SE-24-02"})).status_code == 409
    assert (await admin.delete(f"/api/groups/{g2.id}")).json() == {"deleted": False}

    r = await admin.put("/api/settings", json={"registration_open": False})
    assert r.json()["prefs"]["registration_open"] is False
    assert (await admin.put("/api/settings", json={"min_student_age": 50, "max_student_age": 20})).status_code == 422

    stats = (await admin.get("/api/stats")).json()
    assert stats["students"] == 2 and stats["documents"]["complete"] == 2 and len(stats["per_day"]) == 30

    x = await admin.get("/api/export/students.xlsx?lang=uz")
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["B1"].value == "Familiya" and ws.max_row == 3
    assert ws["N2"].value == "32103051234567"  # PINFL stays text

    log = (await admin.get("/api/audit?action=settings")).json()
    assert log["total"] == 1 and log["items"][0]["details"]["registration_open"] == [True, False]


async def test_spa_and_security_headers(web, h):
    r = await web.client().get("/students/5")
    assert r.status_code == 200 and "text/html" in r.headers["content-type"]
    csp = r.headers["content-security-policy"]
    assert "frame-ancestors 'none'" in csp and "google" not in csp
