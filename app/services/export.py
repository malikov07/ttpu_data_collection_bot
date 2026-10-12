"""Excel export of students, shared by the bot and the website."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.db.models import CertStatus, DocumentKind, Student
from app.i18n import t
from app.services import certificates as certs
from app.services.validators import age_on, format_phone

HEADERS = {
    "en": ["#", "Surname", "Name", "Patronymic", "Group", "Date of birth", "Age", "Gender", "Phone", "Telegram",
           "Document", "Document no.", "Valid until", "PINFL", "Nationality", "Photo", "CV", "Permanent address",
           "Temporary address", "School certificate", "Form 086", "Certificates", "Registered", "Updated"],
    "uz": ["№", "Familiya", "Ism", "Otasining ismi", "Guruh", "Tugʻilgan sana", "Yosh", "Jins", "Telefon", "Telegram",
           "Hujjat", "Hujjat raqami", "Amal qilish muddati", "JShShIR", "Fuqarolik", "Rasm", "Rezyume", "Doimiy manzil",
           "Vaqtinchalik manzil", "Shahodatnoma", "086 shakl", "Sertifikatlar", "Roʻyxatdan oʻtgan", "Yangilangan"],
    "ru": ["№", "Фамилия", "Имя", "Отчество", "Группа", "Дата рождения", "Возраст", "Пол", "Телефон", "Telegram",
           "Документ", "Номер документа", "Действителен до", "ПИНФЛ", "Гражданство", "Фото", "Резюме", "Постоянная прописка",
           "Временная прописка", "Аттестат", "Форма 086", "Сертификаты", "Зарегистрирован", "Обновлён"],
}
WORDS = {
    "en": {"male": "Male", "female": "Female", "passport": "Passport", "id_card": "ID card", "yes": "yes", "no": "no"},
    "uz": {"male": "Erkak", "female": "Ayol", "passport": "Pasport", "id_card": "ID karta", "yes": "bor", "no": "yoʻq"},
    "ru": {"male": "Мужской", "female": "Женский", "passport": "Паспорт", "id_card": "ID-карта", "yes": "есть", "no": "нет"},
}
WIDTHS = [5, 18, 16, 20, 11, 13, 6, 10, 18, 18, 11, 14, 13, 17, 11, 7, 7, 11, 11, 11, 9, 30, 17, 17]
# Document columns after the nationality, in order.
DOCUMENT_COLUMNS = (
    DocumentKind.PHOTO, DocumentKind.CV, DocumentKind.ADDRESS, DocumentKind.TEMP_ADDRESS, DocumentKind.SCHOOL_CERT, DocumentKind.MEDICAL,
)


def certificates_cell(s: Student, lang: str) -> str:
    """Accepted certificates, then those still under review; rejected ones are left out."""
    lines = [certs.label(lang, c) for c in s.certificates if c.status == CertStatus.APPROVED]
    lines += [
        f"{certs.label(lang, c)} ({t(lang, 'cert.status.pending')})" for c in s.certificates if c.status == CertStatus.PENDING
    ]
    return "\n".join(lines)


def students_xlsx(students: list[Student], lang: str = "en") -> bytes:
    lang = lang if lang in HEADERS else "en"
    words = WORDS[lang]
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(HEADERS[lang])
    today = date.today()
    for i, s in enumerate(sorted(students, key=lambda s: (s.group.name, s.full_name.casefold())), start=1):
        has = lambda kind: words["yes"] if s.document(kind) else words["no"]  # noqa: E731
        ws.append([
            i, s.last_name, s.first_name, s.middle_name or "", s.group.name, s.birth_date,
            age_on(s.birth_date, today), words[s.gender.value], format_phone(s.phone),
            f"@{s.user.username}" if s.user and s.user.username else str(s.telegram_id),
            words[s.doc_type.value] if s.doc_type else "", s.doc_number or "", s.doc_expiry,
            s.pinfl or "", s.nationality or "", *(has(kind) for kind in DOCUMENT_COLUMNS),
            certificates_cell(s, lang), s.created_at.replace(tzinfo=None), s.updated_at.replace(tzinfo=None),
        ])
    fill = PatternFill("solid", fgColor="1A4D96")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        row[5].number_format = row[12].number_format = "DD.MM.YYYY"
        row[21].alignment = Alignment(wrap_text=True, vertical="top")
        row[22].number_format = row[23].number_format = "DD.MM.YYYY HH:MM"
        # Long digit strings must stay text (Excel would turn PINFLs into 3.2E+13).
        row[13].number_format = "@"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
